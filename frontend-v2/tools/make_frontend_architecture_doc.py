from pathlib import Path
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
OUT_DIR.mkdir(exist_ok=True)
DIAGRAM_PATH = OUT_DIR / "roadguard_frontend_architecture.png"
DOCX_PATH = OUT_DIR / "RoadGuard_AI_Frontend_Architecture.docx"


COLORS = {
    "ink": "#172033",
    "muted": "#536078",
    "line": "#C9D3E5",
    "panel": "#F7FAFF",
    "blue": "#2563EB",
    "blue_dark": "#1746A2",
    "green": "#16A34A",
    "amber": "#D97706",
    "red": "#DC2626",
    "slate": "#334155",
    "white": "#FFFFFF",
    "future": "#EEF6FF",
}


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def draw_rounded(draw, xy, fill, outline=COLORS["line"], radius=18, width=2):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)


def draw_wrapped(draw, text, box, fill, size=28, bold=False, align="center", line_gap=6):
    x1, y1, x2, y2 = box
    max_width = x2 - x1 - 30
    f = font(size, bold)
    words = text.split()
    lines = []
    current = ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=f)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    line_height = f.size + line_gap
    total_h = len(lines) * line_height - line_gap
    y = y1 + ((y2 - y1 - total_h) / 2)
    for line in lines:
        w = draw.textbbox((0, 0), line, font=f)[2]
        if align == "left":
            x = x1 + 18
        else:
            x = x1 + ((x2 - x1 - w) / 2)
        draw.text((x, y), line, fill=fill, font=f)
        y += line_height


def arrow(draw, start, end, color=COLORS["slate"], width=5):
    draw.line([start, end], fill=color, width=width)
    sx, sy = start
    ex, ey = end
    if abs(ex - sx) >= abs(ey - sy):
        direction = 1 if ex > sx else -1
        pts = [(ex, ey), (ex - direction * 18, ey - 10), (ex - direction * 18, ey + 10)]
    else:
        direction = 1 if ey > sy else -1
        pts = [(ex, ey), (ex - 10, ey - direction * 18), (ex + 10, ey - direction * 18)]
    draw.polygon(pts, fill=color)


def create_diagram():
    img = Image.new("RGB", (2400, 1500), COLORS["white"])
    draw = ImageDraw.Draw(img)

    draw.text((80, 50), "RoadGuard AI Frontend Architecture", fill=COLORS["ink"], font=font(52, True))
    draw.text(
        (82, 116),
        "Current React/Vite implementation with client-side state, maps, mock detection, report workflows, and backend-ready boundaries",
        fill=COLORS["muted"],
        font=font(26),
    )

    # Outer groups
    groups = [
        ((70, 190, 560, 1325), "User Interface Layer", "#F8FAFC"),
        ((620, 190, 1190, 1325), "Routing and Shell", "#F7FBFF"),
        ((1250, 190, 1805, 1325), "Client State and Logic", "#FAFBF7"),
        ((1860, 190, 2330, 1325), "Browser and External Services", "#FFF9F2"),
    ]
    for xy, title, fill in groups:
        draw_rounded(draw, xy, fill=fill, outline="#D6DEED", radius=24, width=3)
        draw.text((xy[0] + 28, xy[1] + 24), title, fill=COLORS["ink"], font=font(30, True))

    def card(x, y, w, h, title, subtitle, fill="#FFFFFF", accent=COLORS["blue"]):
        draw_rounded(draw, (x, y, x + w, y + h), fill=fill, outline="#BFD0E8", radius=18, width=2)
        draw.rectangle((x, y, x + 10, y + h), fill=accent)
        draw.text((x + 28, y + 20), title, fill=COLORS["ink"], font=font(25, True))
        wrapped = textwrap.wrap(subtitle, width=32)
        yy = y + 58
        for line in wrapped[:4]:
            draw.text((x + 28, yy), line, fill=COLORS["muted"], font=font(21))
            yy += 29

    # UI cards
    card(110, 280, 405, 130, "Public Pages", "Landing, Login, Signup through AuthLayout and Navbar", accent=COLORS["blue"])
    card(110, 455, 405, 150, "Dashboard Pages", "Dashboard, Upload, Detection Result, Reports, Map View, Profile", accent=COLORS["green"])
    card(110, 650, 405, 145, "Reusable UI", "Hero, Features, HowItWorks, CTA, Stats, Reveal, Footer", accent=COLORS["amber"])
    card(110, 840, 405, 145, "Design System", "Tailwind CSS theme, lucide-react icons, Framer Motion transitions", accent=COLORS["slate"])
    card(110, 1030, 405, 150, "Map Interface", "React Leaflet maps, CircleMarkers, popups, city search, GPS recentering", accent=COLORS["red"])

    # Routing cards
    card(660, 300, 490, 140, "main.jsx", "React root, BrowserRouter, AuthProvider, DetectionProvider, ReportsProvider", accent=COLORS["blue_dark"])
    card(660, 500, 490, 150, "App.jsx Routes", "Public routes plus protected /app route branch and fallback redirects", accent=COLORS["blue"])
    card(660, 705, 490, 140, "ProtectedRoute", "Allows authenticated users into the app shell and redirects others", accent=COLORS["red"])
    card(660, 905, 490, 165, "DashboardLayout", "Responsive sidebar, header, nav links, logout action, nested Outlet", accent=COLORS["green"])

    # State cards
    card(1290, 280, 470, 135, "AuthContext", "Mock login/signup/profile/logout, persisted as roadguard_user", accent=COLORS["blue"])
    card(1290, 455, 470, 150, "DetectionContext", "Uploaded image URL, mock YOLO-style analysis result, clear/set detection", accent=COLORS["red"])
    card(1290, 650, 470, 155, "ReportsContext", "Draft and submitted reports, save/submit/delete, persisted as roadguard_user_reports", accent=COLORS["green"])
    card(1290, 850, 470, 145, "Mock Data", "Demo reports, severity colors, dashboard distributions, mockDetectionResult", accent=COLORS["amber"])
    card(1290, 1040, 470, 155, "Report Utilities", "generateReportId and browser print window PDF export", accent=COLORS["slate"])

    # External/browser cards
    card(1900, 290, 385, 120, "localStorage", "Current persistence for user and reports", accent=COLORS["blue"])
    card(1900, 455, 385, 130, "Browser APIs", "File input, object URLs, geolocation, print window", accent=COLORS["green"])
    card(1900, 635, 385, 140, "OpenStreetMap Tiles", "Used by React Leaflet through TileLayer", accent=COLORS["amber"])
    card(1900, 830, 385, 150, "Future Backend API", "Auth, image storage, ML inference, reports, PostgreSQL plus PostGIS", fill=COLORS["future"], accent=COLORS["red"])

    # Arrows
    arrow(draw, (515, 520), (660, 570), COLORS["slate"])
    arrow(draw, (1150, 370), (1290, 350), COLORS["slate"])
    arrow(draw, (1150, 570), (1290, 530), COLORS["slate"])
    arrow(draw, (1150, 980), (1290, 725), COLORS["slate"])
    arrow(draw, (1760, 345), (1900, 350), COLORS["slate"])
    arrow(draw, (1760, 530), (1900, 520), COLORS["slate"])
    arrow(draw, (1760, 930), (1900, 905), COLORS["red"], width=4)

    # Flow strip
    y = 1365
    draw.text((85, y), "Primary user flow:", fill=COLORS["ink"], font=font(26, True))
    steps = [
        "Login/Signup",
        "Protected App Shell",
        "Upload Photo + GPS",
        "Mock Analysis",
        "Save Draft or Submit",
        "Reports + Map + PDF",
    ]
    x = 310
    for i, step in enumerate(steps):
        draw_rounded(draw, (x, y - 12, x + 265, y + 55), fill="#FFFFFF", outline="#BFD0E8", radius=14, width=2)
        draw_wrapped(draw, step, (x, y - 12, x + 265, y + 55), COLORS["ink"], size=20, bold=True)
        if i < len(steps) - 1:
            arrow(draw, (x + 265, y + 22), (x + 310, y + 22), COLORS["blue"], width=4)
        x += 315

    img.save(DIAGRAM_PATH, quality=95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_border(cell, color="D9D9D9"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def format_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(4)
                for run in p.runs:
                    run.font.name = "Aptos"
                    run.font.size = Pt(9.5)
                    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
            if row_i == 0:
                set_cell_shading(cell, "#1F2937")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.bold = True


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    format_table(table)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
    return table


def set_doc_defaults(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(32, 41, 58)
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        p_pr = style._element.get_or_add_pPr()
        for existing in p_pr.findall(qn("w:pBdr")):
            p_pr.remove(existing)


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        run = p.add_run(bold_lead)
        run.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    set_doc_defaults(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(2)
    title.add_run("RoadGuard AI Frontend Architecture")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(6)
    subtitle.add_run("Architecture diagram and implementation notes for the current React frontend")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.add_run(
        "This document explains how the RoadGuard AI frontend is structured today. "
        "It covers routing, page composition, shared state, browser storage, maps, report actions, "
        "and the realistic backend boundary for adding PostgreSQL with PostGIS later."
    )

    doc.add_picture(str(DIAGRAM_PATH), width=Inches(9.7))

    doc.add_heading("Current Frontend Layers", level=1)
    add_table(
        doc,
        ["Layer", "Implementation", "Responsibility"],
        [
            ["Application boot", "src/main.jsx", "Mounts React, enables BrowserRouter, and wraps App with Auth, Detection, and Reports providers."],
            ["Routing", "src/App.jsx", "Defines public routes, protected app routes, nested dashboard routes, and redirects."],
            ["Layout shell", "DashboardLayout and AuthLayout", "Provides dashboard navigation, responsive sidebar, header controls, auth page framing, and nested content outlet."],
            ["Feature pages", "Dashboard, Upload, DetectionResult, Reports, MapView, Profile", "Implements the citizen reporting workflow, map view, profile management, and report browsing."],
            ["Shared UI", "Hero, Navbar, Footer, CTA, Features, HowItWorks, Stats, Reveal", "Supports the landing experience and repeated presentation sections."],
            ["Styling and motion", "Tailwind CSS, Framer Motion, lucide-react", "Controls the visual system, responsive layout, animation, and iconography."],
        ],
        widths=[Cm(4.2), Cm(6.0), Cm(15.7)],
    )

    doc.add_heading("Primary Data Flow", level=1)
    add_para(
        doc,
        "Current flow: A user signs in through the mock AuthContext, enters the protected app shell, uploads or captures a road image, attaches GPS through the browser geolocation API, runs the mock analysis flow, then saves a draft or submits the report through ReportsContext."
    )
    add_table(
        doc,
        ["Step", "Frontend module", "State or output"],
        [
            ["Authenticate", "Login.jsx, Signup.jsx, AuthContext", "Creates a mock user and persists roadguard_user in localStorage."],
            ["Capture evidence", "Upload.jsx", "Stores a browser object URL for the selected image and requests GPS coordinates."],
            ["Analyze", "DetectionContext, mockData", "Combines mockDetectionResult with image, report id, GPS, severity, score, and priority."],
            ["Review result", "DetectionResult.jsx", "Shows damage overlay, metrics, map pin, and actions for draft, submit, and PDF download."],
            ["Persist report", "ReportsContext", "Stores drafts and submitted reports in roadguard_user_reports in localStorage."],
            ["Visualize", "Reports.jsx, MapView.jsx, Dashboard.jsx", "Merges user reports with mock reports for tables, charts, filters, and Leaflet markers."],
        ],
        widths=[Cm(4.1), Cm(7.0), Cm(14.8)],
    )

    doc.add_heading("State and Persistence", level=1)
    add_table(
        doc,
        ["Area", "Current frontend source", "Current persistence", "Backend ready replacement"],
        [
            ["User session", "AuthContext", "localStorage key roadguard_user", "Real auth session from backend or Supabase Auth."],
            ["Detection result", "DetectionContext", "React state only", "ML inference API response linked to uploaded image and report id."],
            ["Reports", "ReportsContext", "localStorage key roadguard_user_reports", "Reports table in PostgreSQL with PostGIS geometry for coordinates."],
            ["Images", "Upload.jsx object URL", "Temporary browser object URL", "Object storage path plus metadata stored on report record."],
            ["Map data", "mockReports plus user reports", "Bundled mock data and localStorage", "API query with geospatial filtering, severity filters, and pagination."],
            ["PDF export", "reportUtils.downloadReportPDF", "Browser print window", "Client generated PDF or backend generated official report document."],
        ],
        widths=[Cm(4.1), Cm(5.8), Cm(6.0), Cm(10.0)],
    )

    doc.add_heading("Backend Integration Boundary", level=1)
    add_para(
        doc,
        "The frontend can keep its current route and page structure when a backend is added. The main change should be replacing context internals with service calls while preserving the page-level API used by the components."
    )
    add_table(
        doc,
        ["Frontend contract", "Future API or service", "Why it matters"],
        [
            ["login, signup, logout, updateProfile", "Authentication service", "Moves mock auth out of localStorage and enables real identity, roles, and session expiry."],
            ["analyzeImage(imageUrl, gps, reportId)", "Image upload plus detection inference endpoint", "Connects captured road evidence to ML output and permanent storage."],
            ["saveDraft, submitReport, deleteReport", "Reports API", "Allows report lifecycle tracking, moderation, authority review, and multi-device access."],
            ["MapView report query", "PostgreSQL/PostGIS search endpoint", "Supports nearby potholes, bounding box queries, city filters, and severity-based map layers."],
            ["downloadReportPDF", "Report export endpoint or client PDF renderer", "Produces consistent downloadable documents for submitted complaints."],
        ],
        widths=[Cm(6.0), Cm(7.4), Cm(12.5)],
    )

    doc.add_heading("Recommended Frontend Direction", level=1)
    add_para(
        doc,
        "Keep the current React Router structure and provider boundaries because they already match the product workflow. As the project grows, introduce a small services layer between contexts and backend APIs, then migrate localStorage data to PostgreSQL/PostGIS through authenticated endpoints."
    )
    add_para(
        doc,
        "For the next version, the most useful frontend change is to create dedicated API modules for auth, reports, detection, uploads, and map queries. That will let the pages stay stable while persistence moves from the browser to the backend."
    )

    doc.core_properties.title = "RoadGuard AI Frontend Architecture"
    doc.core_properties.subject = "Frontend architecture diagram and implementation notes"
    doc.core_properties.author = "RoadGuard AI"
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    create_diagram()
    create_docx()
    print(DOCX_PATH)
